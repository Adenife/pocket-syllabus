import tempfile
import unittest
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from docx import Document
from app.ingestion.pipeline import parse_document
from helpers import context


def pdf_fixture(path, text='Light makes photosynthesis possible.', blank=False):
    writer=PdfWriter()
    page=writer.add_blank_page(width=600,height=800)
    if not blank:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 50 700 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
    writer.write(path)

class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)
    def test_pdf_upload_index_page_citation(self):
        path=self.root/'book.pdf'
        pdf_fixture(path)
        ctx=context(self.root/'library')
        result=ctx.documents.upload('local',path.name,path.read_bytes())
        ctx.worker.run_once()
        doc=ctx.documents.require(result['document']['id'],'local')
        self.assertEqual(doc['status'],'ready',doc['error'])
        results=ctx.search.search('photosynthesis')
        self.assertEqual(results[0].page,1)
    def test_scanned_pdf_reports_ocr_requirement(self):
        path=self.root/'scan.pdf'
        pdf_fixture(path,blank=True)
        ctx=context(self.root/'library')
        result=ctx.documents.upload('local',path.name,path.read_bytes())
        ctx.worker.run_once()
        doc=ctx.documents.require(result['document']['id'],'local')
        self.assertEqual(doc['status'],'failed')
        self.assertIn('OCR',doc['error'])
        self.assertTrue(doc['warnings'])
    def test_docx_sections_and_tables(self):
        path=self.root/'book.docx'
        document=Document()
        document.add_heading('Physics',level=1)
        document.add_heading('Motion',level=2)
        document.add_paragraph('Motion is a change in position.')
        table=document.add_table(rows=1,cols=2)
        table.cell(0,0).text='Speed'
        table.cell(0,1).text='Distance / time'
        document.save(path)
        warnings=[]
        blocks=list(parse_document(path,path.name,warnings))
        self.assertEqual(blocks[0].section,'Physics / Motion')
        self.assertIn('Speed | Distance / time',blocks[1].text)
        ctx=context(self.root/'library')
        result=ctx.documents.upload('local',path.name,path.read_bytes())
        ctx.worker.run_once()
        self.assertEqual(ctx.documents.require(result['document']['id'],'local')['status'],'ready')
    def test_txt_paragraphs_keep_line_numbers(self):
        path=self.root/'notes.txt'
        path.write_text('First line\nnext line\n\nSecond paragraph')
        blocks=list(parse_document(path,path.name,[]))
        self.assertEqual([b.line for b in blocks],[1,4])
